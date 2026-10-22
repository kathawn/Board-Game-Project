#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tien ich trinh bay Word dung chung cho fill_brd.py / fill_srs.py / build_sdd.py.
Dong bo phong cach voi build_checklist.py: font Calibri, tieu de mau Navy #1B365D,
bang ke vien tinh te, nen xen ke (zebra).
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1B, 0x36, 0x5D)
NAVY_HEX = "1B365D"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ZEBRA_HEX = "EEF2F7"
GREY = RGBColor(0x44, 0x44, 0x44)
GREY_HEX = "5B5B5B"
GRID_BORDER_HEX = "B9C4D4"


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color_hex=GRID_BORDER_HEX, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color_hex)
        borders.append(el)
    tcPr.append(borders)


def format_run(run, size=10.5, bold=False, italic=False, color=None, font_name="Calibri"):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def new_document():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Calibri")
    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    for lvl, size, color in ((1, 16, NAVY), (2, 13, NAVY), (3, 11.5, NAVY)):
        hstyle = doc.styles[f"Heading {lvl}"]
        hstyle.font.name = "Calibri"
        hstyle.font.size = Pt(size)
        hstyle.font.bold = True
        hstyle.font.color.rgb = color
        hstyle.font.italic = False
    return doc


def title_page(doc, main_title, subtitle, project_name, author, doc_code):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    run = p.add_run(main_title)
    format_run(run, size=24, bold=True, color=NAVY)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(f"({subtitle})")
    format_run(run2, size=13, italic=True, color=GREY)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(20)
    run3 = p3.add_run(project_name)
    format_run(run3, size=18, bold=True, color=RGBColor(0x44, 0x44, 0x44))

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(80)
    run4 = p4.add_run(f"Mã tài liệu: {doc_code}  |  Phiên bản: 1.0")
    format_run(run4, size=10.5, color=GREY)

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run5 = p5.add_run(f"Người soạn thảo: {author}")
    format_run(run5, size=10.5, color=GREY)
    doc.add_page_break()


def version_history(doc, rows):
    heading(doc, "LỊCH SỬ PHIÊN BẢN", level=1)
    table(doc, ["Phiên bản", "Ngày", "Người sửa", "Mô tả thay đổi"], rows,
          col_widths_cm=[2.2, 2.6, 3.2, 8.5])
    doc.add_page_break()


def toc_placeholder(doc):
    heading(doc, "MỤC LỤC", level=1)
    p = doc.add_paragraph()
    run = p.add_run(
        "(Mục lục tự động: trong Word, chọn thẻ References > Table of Contents, hoặc nếu đã có "
        "trường Mục lục thì bấm chuột phải > Update Field > Update entire table.)"
    )
    format_run(run, size=9.5, italic=True, color=GREY)
    doc.add_page_break()


def heading(doc, text, level=1):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    sizes = {1: 16, 2: 13, 3: 11.5}
    format_run(run, size=sizes.get(level, 11), bold=True, color=NAVY)
    return h


def para(doc, text, bold=False, italic=False, size=10.5, color=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    format_run(run, size=size, bold=bold, italic=italic, color=color or RGBColor(0x22, 0x22, 0x22))
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        format_run(run, size=10.5, color=RGBColor(0x22, 0x22, 0x22))


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        format_run(run, size=10.5, color=RGBColor(0x22, 0x22, 0x22))


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    format_run(run, size=9.5, italic=True, bold=True, color=GREY)
    p.paragraph_format.space_after = Pt(12)
    return p


def picture(doc, path, width_cm=15, cap=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    if cap:
        caption(doc, cap)


def table(doc, headers, rows, col_widths_cm=None, first_col_bold=False):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    for idx, text in enumerate(headers):
        cell = t.rows[0].cells[idx]
        cell.text = ""
        set_cell_shading(cell, NAVY_HEX)
        set_cell_borders(cell, color_hex=NAVY_HEX)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        format_run(run, size=10, bold=True, color=WHITE)

    for r_idx, row_vals in enumerate(rows, start=1):
        zebra = r_idx % 2 == 0
        bg = ZEBRA_HEX if zebra else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = ""
            set_cell_shading(cell, bg)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            bold = first_col_bold and c_idx == 0
            format_run(run, size=9.5, bold=bold, color=NAVY if bold else RGBColor(0x22, 0x22, 0x22))

    if col_widths_cm:
        for row in t.rows:
            for idx, w in enumerate(col_widths_cm):
                row.cells[idx].width = Cm(w)
        for idx, w in enumerate(col_widths_cm):
            t.columns[idx].width = Cm(w)
    doc.add_paragraph()
    return t


def code_block(doc, text):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F3F4F6")
    pPr.append(shd)
    p.paragraph_format.space_after = Pt(2)
    for i, line in enumerate(text.strip("\n").split("\n")):
        run = p.add_run((line or " "))
        format_run(run, size=9.5, color=RGBColor(0x1B, 0x1B, 0x1B), font_name="Consolas")
        if i < len(text.strip("\n").split("\n")) - 1:
            run.add_break()
    doc.add_paragraph()
    return p


def note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    format_run(run, size=9.5, italic=True, color=GREY)
    return p


def sign_off_table(doc, approvers):
    """approvers: list of (ten, vai_tro, trach_nhiem)."""
    rows = [[name, role, "", ""] for name, role, _desc in approvers]
    table(doc, ["Họ tên / Đại diện", "Vai trò", "Chữ ký", "Ngày ký"],
          rows, col_widths_cm=[5.0, 5.5, 3.0, 2.5])
