import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import extract_spec


class TestExtractFromMarkdown(unittest.TestCase):
    def setUp(self):
        self.md_path = Path(self.enterContext(
            __import__("tempfile").TemporaryDirectory()
        )) / "srs.md"
        self.md_path.write_text(
            "# 1. Giới thiệu\n"
            "Nội dung giới thiệu chung.\n\n"
            "# 2. UC-04 Đăng nhập\n"
            "Khách hàng nhập email và mật khẩu để đăng nhập hệ thống.\n\n"
            "# 3. UC-07 Đặt cọc\n"
            "Khách hàng đặt cọc trước khi nhận hàng.\n",
            encoding="utf-8",
        )

    def test_finds_section_by_uc_id(self):
        sections = extract_spec.extract(self.md_path, "UC-04")
        self.assertEqual(len(sections), 1)
        self.assertIn("Đăng nhập", sections[0].heading)
        self.assertIn("mật khẩu", sections[0].text)

    def test_finds_section_by_keyword_in_body(self):
        sections = extract_spec.extract(self.md_path, "đặt cọc")
        self.assertEqual(len(sections), 1)
        self.assertIn("UC-07", sections[0].heading)

    def test_no_match_raises_nomatcherror(self):
        with self.assertRaises(extract_spec.NoMatchError):
            extract_spec.extract(self.md_path, "khong-ton-tai-xyz")


class TestExtractFromDocx(unittest.TestCase):
    def setUp(self):
        import docx
        tmp = Path(self.enterContext(__import__("tempfile").TemporaryDirectory()))
        self.docx_path = tmp / "srs.docx"
        doc = docx.Document()
        doc.add_heading("1. Giới thiệu", level=1)
        doc.add_paragraph("Nội dung giới thiệu chung.")
        doc.add_heading("3.4.4 UC-04 Đăng nhập", level=2)
        doc.add_paragraph("Khách hàng nhập email và mật khẩu để đăng nhập.")
        doc.save(str(self.docx_path))

    def test_finds_section_by_uc_id(self):
        sections = extract_spec.extract(self.docx_path, "UC-04")
        self.assertEqual(len(sections), 1)
        self.assertIn("Đăng nhập", sections[0].heading)
        self.assertIn("mật khẩu", sections[0].text)

    def test_docx_with_no_headings_falls_back_to_whole_document(self):
        import docx
        tmp = Path(self.enterContext(__import__("tempfile").TemporaryDirectory()))
        flat_path = tmp / "flat.docx"
        doc = docx.Document()
        doc.add_paragraph("Văn bản không có heading nào, chỉ có UC-09 nhắc tới ở đây.")
        doc.save(str(flat_path))
        sections = extract_spec.extract(flat_path, "UC-09")
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0].heading, "(toàn văn)")


class TestExtractFromPdf(unittest.TestCase):
    def test_raises_missing_dependency_when_pypdf_absent(self):
        with mock.patch.dict(sys.modules, {"pypdf": None}):
            with self.assertRaises(extract_spec.MissingDependencyError) as ctx:
                extract_spec.extract_from_pdf(Path("dummy.pdf"), "query")
        self.assertIn("pip install pypdf", str(ctx.exception))


class TestExtractDispatchAndCli(unittest.TestCase):
    def test_unsupported_extension_raises_valueerror(self):
        with self.assertRaises(ValueError):
            extract_spec.extract(Path("file.txt"), "query")

    def test_main_prints_no_match_message_and_exits_nonzero(self):
        tmp = Path(self.enterContext(__import__("tempfile").TemporaryDirectory()))
        md_path = tmp / "srs.md"
        md_path.write_text("# Mục 1\nNội dung A.\n", encoding="utf-8")
        with self.assertRaises(SystemExit) as ctx:
            extract_spec.main([str(md_path), "khong-ton-tai"])
        self.assertNotEqual(ctx.exception.code, 0)


if __name__ == "__main__":
    unittest.main()
