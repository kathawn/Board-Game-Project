"""Extract the relevant section(s) of an SRS document (.md, .docx, .pdf).

Given a keyword or UC-id, pulls out just the matching section(s) of a
spec document so callers don't need to read the whole document into
context.
"""

import argparse
import dataclasses
import sys
from pathlib import Path


@dataclasses.dataclass
class Section:
    heading: str
    text: str
    location: str


class MissingDependencyError(RuntimeError):
    pass


class NoMatchError(RuntimeError):
    pass


def _matches(query: str, heading: str, text: str) -> bool:
    query_lower = query.lower()
    return query_lower in heading.lower() or query_lower in text.lower()


def extract_from_md(path: Path, query: str) -> list[Section]:
    content = path.read_text(encoding="utf-8")
    lines = content.splitlines()

    blocks: list[tuple[str, list[str]]] = []
    current_heading = None
    current_lines: list[str] = []
    for line in lines:
        if line.startswith("#"):
            if current_heading is not None:
                blocks.append((current_heading, current_lines))
            current_heading = line.lstrip("#").strip()
            current_lines = []
        else:
            current_lines.append(line)
    if current_heading is not None:
        blocks.append((current_heading, current_lines))

    sections: list[Section] = []
    for heading, body_lines in blocks:
        text = "\n".join(body_lines).strip()
        if _matches(query, heading, text):
            sections.append(Section(heading=heading, text=text, location=heading))

    if not sections:
        raise NoMatchError(query)
    return sections


def extract_from_docx(path: Path, query: str) -> list[Section]:
    import docx

    doc = docx.Document(str(path))

    blocks: list[tuple[str, list[str]]] = []
    current_heading = None
    current_lines: list[str] = []
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.startswith("Heading") or style_name == "Title":
            if current_heading is not None:
                blocks.append((current_heading, current_lines))
            current_heading = paragraph.text
            current_lines = []
        else:
            current_lines.append(paragraph.text)
    if current_heading is not None:
        blocks.append((current_heading, current_lines))

    if not blocks:
        whole_text = "\n".join(p.text for p in doc.paragraphs).strip()
        blocks = [("(toàn văn)", whole_text.splitlines())]

    sections: list[Section] = []
    for heading, body_lines in blocks:
        text = "\n".join(body_lines).strip()
        if _matches(query, heading, text):
            sections.append(Section(heading=heading, text=text, location=heading))

    if not sections:
        raise NoMatchError(query)
    return sections


def extract_from_pdf(path: Path, query: str) -> list[Section]:
    try:
        import pypdf
    except ImportError:
        raise MissingDependencyError(
            "Thiếu thư viện pypdf — chạy: pip install pypdf"
        )

    reader = pypdf.PdfReader(path)

    sections: list[Section] = []
    for page_number, page in enumerate(reader.pages, start=1):
        heading = f"Trang {page_number}"
        text = (page.extract_text() or "").strip()
        if _matches(query, heading, text):
            sections.append(Section(heading=heading, text=text, location=heading))

    if not sections:
        raise NoMatchError(query)
    return sections


def extract(path: Path, query: str) -> list[Section]:
    suffix = path.suffix.lower()
    if suffix == ".md":
        return extract_from_md(path, query)
    if suffix == ".docx":
        return extract_from_docx(path, query)
    if suffix == ".pdf":
        return extract_from_pdf(path, query)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description="Extract matching section(s) from an SRS document."
    )
    parser.add_argument("file", help="Path to the SRS document (.md, .docx, .pdf)")
    parser.add_argument("query", help="Keyword or UC-id to search for")
    parser.add_argument(
        "--max-results",
        type=int,
        default=3,
        help="Maximum number of matching sections to print (default: 3)",
    )
    args = parser.parse_args(argv)

    try:
        sections = extract(Path(args.file), args.query)
    except NoMatchError:
        print(f"No matching section found for query: {args.query}", file=sys.stderr)
        sys.exit(1)
    except MissingDependencyError as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)

    for section in sections[: args.max_results]:
        print(f"=== {section.heading} ({section.location}) ===")
        print(section.text)
        print()


if __name__ == "__main__":
    main()
